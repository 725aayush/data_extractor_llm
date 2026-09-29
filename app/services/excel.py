import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _sheet_name(name: object, used: set[str], fallback: str) -> str:
    base = re.sub(r"[\[\]*:/\\?]", "", str(name or "")).strip() or fallback
    base = base[:31]
    candidate, counter = base, 2
    while candidate in used:
        suffix = f"_{counter}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _table_frame(table: dict) -> pd.DataFrame:
    columns = [str(c.get("name", "")).strip() for c in table.get("columns", []) if isinstance(c, dict) and str(c.get("name", "")).strip()]
    rows = []
    for row in table.get("rows", []):
        if not isinstance(row, dict) or not isinstance(row.get("cells"), dict):
            continue
        cells = row["cells"]
        for key in cells:
            if str(key) not in columns:
                columns.append(str(key))
        rows.append({**{key: cells.get(key, "") for key in columns}, "_page": row.get("page", ""), "_confidence": row.get("confidence", "")})
    return pd.DataFrame(rows, columns=columns + ["_page", "_confidence"])


def create_excel(data: dict, output_path: Path) -> None:
    used: set[str] = set()
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        info = []
        for section in ("document", "metadata", "seller", "buyer"):
            for field, value in data.get(section, {}).items() if isinstance(data.get(section), dict) else []:
                info.append({"Section": section.title(), "Field": field, "Value": value})
        if info:
            pd.DataFrame(info).to_excel(writer, sheet_name="Invoice Info", index=False)
            used.add("Invoice Info")
        for number, table in enumerate(data.get("tables", []), start=1):
            if isinstance(table, dict):
                frame = _table_frame(table)
                if not frame.empty:
                    frame.to_excel(writer, sheet_name=_sheet_name(table.get("title") or table.get("table_id"), used, f"Table {number}"), index=False)
        for source, title in (("taxes", "Taxes"), ("other_information", "Other Info"), ("pages", "Pages")):
            values = data.get(source, [])
            if values:
                pd.DataFrame(values).to_excel(writer, sheet_name=title, index=False)
                used.add(title)
        totals = data.get("totals", {})
        if isinstance(totals, dict) and totals:
            pd.DataFrame([{"Field": key, "Value": value} for key, value in totals.items()]).to_excel(writer, sheet_name="Totals", index=False)
            used.add("Totals")
        validation = data.get("validation", {})
        validation_rows = [{"Field": "Needs Review", "Value": validation.get("needs_review", False)}]
        validation_rows += [{"Field": f"Issue {i}", "Value": json.dumps(issue, ensure_ascii=False)} for i, issue in enumerate(validation.get("issues", []), 1)]
        pd.DataFrame(validation_rows).to_excel(writer, sheet_name="Validation", index=False)
    _format_excel(output_path)


def _format_excel(path: Path) -> None:
    workbook = load_workbook(path)
    fill, font = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill, cell.font = fill, font
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max((len(str(c.value or "")) for c in column), default=0) + 3, 60)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)
