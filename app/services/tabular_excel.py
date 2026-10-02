"""Business-readable Excel export that mirrors the logical invoice structure."""

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


STANDARD_ITEM_COLUMNS = (
    "Item No.", "Description", "HSN", "Quantity", "Rate", "Unit",
    "Amount", "Taxable Value", "GST Rate", "GST Amount",
)


def _label(value: object) -> str:
    return str(value).replace("_", " ").strip().title()


def _field_rows(*sections: tuple[str, dict]) -> list[dict]:
    rows: list[dict] = []
    for title, values in sections:
        if not isinstance(values, dict):
            continue
        for field, value in values.items():
            if isinstance(value, (str, int, float, bool)) or value is None:
                rows.append({"Section": title, "Field": _label(field), "Value": value or ""})
    return rows or [{"Section": "Details", "Field": "Status", "Value": "No details extracted"}]


def _normalise_column(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def _standard_item_column(value: object) -> str | None:
    column = _normalise_column(value)
    if column in {"item", "item no", "item number", "serial no", "sr no", "line no"}:
        return "Item No."
    if "description" in column or column in {"goods", "product", "item description"}:
        return "Description"
    if column in {"hsn", "sac", "hsn sac", "hsn code"}:
        return "HSN"
    if column in {"qty", "quantity"}:
        return "Quantity"
    if column in {"rate", "unit rate", "unit price", "price"}:
        return "Rate"
    if column in {"unit", "uom"}:
        return "Unit"
    if "taxable" in column and "value" in column:
        return "Taxable Value"
    if ("gst" in column or "igst" in column or "tax" in column) and ("rate" in column or "%" in str(value)):
        return "GST Rate"
    if ("gst" in column or "igst" in column or "tax" in column) and "amount" in column:
        return "GST Amount"
    if column in {"amount", "line amount", "line total", "total amount", "value"}:
        return "Amount"
    return None


def _line_items(data: dict) -> pd.DataFrame:
    records: list[dict] = []
    for table in data.get("tables", []):
        if not isinstance(table, dict):
            continue
        for number, row in enumerate(table.get("rows", []), start=1):
            cells = row.get("cells", {}) if isinstance(row, dict) else {}
            if not isinstance(cells, dict):
                continue
            record = {column: "" for column in STANDARD_ITEM_COLUMNS}
            record["Item No."] = number
            extra: dict[str, object] = {}
            for source_column, value in cells.items():
                destination = _standard_item_column(source_column)
                if destination and not record[destination]:
                    record[destination] = value
                else:
                    extra[_label(source_column)] = value
            record.update(extra)
            records.append(record)
    extra_columns = list(dict.fromkeys(key for row in records for key in row if key not in STANDARD_ITEM_COLUMNS))
    return pd.DataFrame(records, columns=[*STANDARD_ITEM_COLUMNS, *extra_columns])


def _write_details_sheet(writer: pd.ExcelWriter, name: str, *sections: tuple[str, dict]) -> None:
    pd.DataFrame(_field_rows(*sections)).to_excel(writer, sheet_name=name, index=False)


def _write_tax_sheet(writer: pd.ExcelWriter, data: dict) -> None:
    rows = _field_rows(("Totals", data.get("totals", {})))
    for tax in data.get("taxes", []):
        if isinstance(tax, dict):
            rows.extend({"Section": "Tax", "Field": _label(key), "Value": value} for key, value in tax.items())
    pd.DataFrame(rows).to_excel(writer, sheet_name="Tax & Totals", index=False)


def create_tabular_excel(data: dict, output_path: Path) -> None:
    """Create logical invoice sheets with master details and line items."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_details_sheet(writer, "Invoice Details", ("Document", data.get("document", {})), ("Invoice", data.get("metadata", {})))
        _write_details_sheet(writer, "Seller Details", ("Seller", data.get("seller", {})))
        _write_details_sheet(writer, "Buyer Bill-To", ("Buyer / Bill-To", data.get("buyer", {})))
        _write_details_sheet(writer, "Ship-To Details", ("Ship-To", data.get("ship_to", {})), ("Shipping", data.get("shipping", {})))
        _line_items(data).to_excel(writer, sheet_name="Goods Line Items", index=False)
        _write_tax_sheet(writer, data)
        other = _field_rows(("Other Details", data.get("other_information", {})), ("Validation", data.get("validation", {})))
        if isinstance(data.get("other_information"), list):
            other = [{"Section": "Other Details", "Field": item.get("label", ""), "Value": item.get("value", "")} for item in data["other_information"] if isinstance(item, dict)] or other
        pd.DataFrame(other).to_excel(writer, sheet_name="Other Details", index=False)
    _format_business_workbook(output_path)


def _format_business_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill, cell.font, cell.alignment, cell.border = header_fill, white_bold, Alignment(horizontal="center", vertical="center", wrap_text=True), border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
        if sheet.title != "Goods Line Items":
            row = 2
            while row <= sheet.max_row:
                section = sheet.cell(row, 1).value
                end_row = row
                while end_row < sheet.max_row and sheet.cell(end_row + 1, 1).value == section:
                    end_row += 1
                for cell in sheet[row]:
                    cell.fill = section_fill
                    cell.font = Font(bold=True)
                if end_row > row:
                    sheet.merge_cells(start_row=row, start_column=1, end_row=end_row, end_column=1)
                    sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                row = end_row + 1
        for index, column in enumerate(sheet.columns, start=1):
            width = min(max(max((len(str(cell.value or "")) for cell in column), default=0) + 3, 14), 45)
            if sheet.title != "Goods Line Items" and index == 3:
                width = 55
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 28
    workbook.save(path)
